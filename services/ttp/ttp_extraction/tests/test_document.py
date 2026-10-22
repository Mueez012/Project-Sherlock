import io
import json

import pytest

from ttp_extraction.inference.document import DocumentError, clean_text, extract_text

PROSE = ("The actor sent spearphishing emails with a malicious attachment. "
         "A macro then launched PowerShell to download the second stage.")


def make_pdf(pages: list[str]) -> bytes:
    import pymupdf

    pdf = pymupdf.open()
    for text in pages:
        page = pdf.new_page()
        page.insert_textbox(pymupdf.Rect(50, 50, 550, 750), text, fontsize=11)
    return pdf.tobytes()


def test_pdf_with_text_layer_keeps_pages():
    doc = extract_text(make_pdf([PROSE, "Second page. The backdoor beaconed over HTTPS every hour."]), "r.pdf")
    assert doc.pages == 2
    assert "spearphishing" in doc.text and "beaconed" in doc.text
    second = doc.text.index("Second page")
    assert doc.page_of(second) == 2 and doc.page_of(0) == 1


def test_repeated_pdf_headers_are_removed():
    pages = [f"ACME CONFIDENTIAL\n\nPage body {i}: the actor used tool number {i} for discovery."
             for i in range(4)]
    doc = extract_text(make_pdf(pages), "r.pdf")
    assert "ACME CONFIDENTIAL" not in doc.text
    assert any("header" in w for w in doc.warnings)


def test_scanned_pdf_without_text_layer_is_rejected():
    import pymupdf

    pdf = pymupdf.open()
    page = pdf.new_page()
    page.draw_rect(pymupdf.Rect(50, 50, 300, 300), color=(0, 0, 0), fill=(0.5, 0.5, 0.5))   # image-like, no text
    with pytest.raises(DocumentError) as err:
        extract_text(pdf.tobytes(), "scan.pdf")
    assert err.value.code == "no_text_layer"


def test_scanned_pdf_is_read_with_ocr():
    import pymupdf

    from ttp_extraction.inference.document import ocr_available

    if not ocr_available():
        pytest.skip("RapidOCR not installed")
    src = pymupdf.open()
    page = src.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 750), PROSE, fontsize=13)
    scan = pymupdf.open()
    image_page = scan.new_page(width=page.rect.width, height=page.rect.height)
    image_page.insert_image(image_page.rect, pixmap=page.get_pixmap(dpi=150))   # pixels only, no text layer
    doc = extract_text(scan.tobytes(), "scan.pdf")
    assert "spearphishing" in doc.text and "PowerShell" in doc.text
    assert any("OCR" in w for w in doc.warnings)


def test_scanned_pdf_without_ocr_is_rejected():
    import pymupdf

    from ttp_extraction.inference.document import read_pdf

    pdf = pymupdf.open()
    pdf.new_page()
    with pytest.raises(DocumentError) as err:
        read_pdf(pdf.tobytes(), ocr=False)
    assert err.value.code == "no_text_layer"


def test_txt_utf8_and_latin1():
    assert "café" in extract_text("The café server was hit.".encode("utf-8"), "a.txt").text
    doc = extract_text("Le serveur du café a été compromis par l'acteur.".encode("latin-1"), "b.txt")
    assert "café" in doc.text and doc.warnings


def test_html_strips_scripts_and_nav():
    html = f"""<html><head><script>var secret = 1;</script><style>p{{}}</style></head>
    <body><nav>Home | About</nav><article><h1>Report</h1><p>{PROSE}</p></article>
    <footer>Copyright</footer></body></html>"""
    doc = extract_text(html.encode(), "r.html")
    assert "spearphishing" in doc.text
    assert "secret" not in doc.text and "Home | About" not in doc.text and "Copyright" not in doc.text


def test_docx_paragraphs():
    import docx

    d = docx.Document()
    d.add_paragraph(PROSE)
    d.add_paragraph("The operators used scheduled tasks for persistence.")
    buf = io.BytesIO()
    d.save(buf)
    doc = extract_text(buf.getvalue(), "r.docx")
    assert "scheduled tasks" in doc.text and doc.text.count("\n\n") >= 1


def test_stix_extracts_descriptions_but_not_attack_pattern_ids():
    bundle = {"type": "bundle", "id": "bundle--1", "objects": [
        {"type": "report", "id": "report--1", "name": "Op Test", "description": PROSE},
        {"type": "malware", "id": "malware--1", "name": "EvilLoader", "description": "A loader written in C."},
        {"type": "attack-pattern", "id": "attack-pattern--1", "name": "PowerShell",
         "description": "SHOULD NOT BE READ",
         "external_references": [{"source_name": "mitre-attack", "external_id": "T1059.001"}]},
    ]}
    doc = extract_text(json.dumps(bundle).encode(), "b.json")
    assert "spearphishing" in doc.text and "EvilLoader" in doc.text
    assert "SHOULD NOT BE READ" not in doc.text and "T1059" not in doc.text
    assert doc.declared_ttps == ["T1059.001"]


def test_json_that_is_not_stix_is_rejected():
    with pytest.raises(DocumentError) as err:
        extract_text(b'{"hello": "world"}', "x.json")
    assert err.value.code == "unsupported_file"


def test_cleaning_keeps_defanged_iocs_and_fixes_hyphen_breaks():
    text = clean_text("The mal-\nware contacted hxxp://evil[.]com and “quoted”  text.")
    assert "malware" in text and "hxxp://evil[.]com" in text and '"quoted" text' in text


def test_empty_text_is_rejected():
    with pytest.raises(DocumentError) as err:
        extract_text(b"   \n\n  ", "e.txt")
    assert err.value.code == "no_text"
