"""Uploaded file -> clean text (+ page map, warnings) for the TTP model.

Supported: PDF (PyMuPDF; scanned pages are OCR'd with RapidOCR when it is
installed), TXT / MD (UTF-8, then charset detection), HTML
(BeautifulSoup main text), DOCX (python-docx) and STIX 2.1 JSON bundles.

The output text uses a blank line between paragraphs; segment.py relies on
that. For PDFs, `page_starts[i]` is the character offset where page i + 1
begins, so every sentence can be traced back to its page.

For STIX bundles the existing attack-pattern ids are NOT put in the text the
model reads (that would let the model "see the answer"); they are returned
separately as `declared_ttps` so the UI can compare.
"""
from __future__ import annotations

import io
import json
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from pathlib import Path

EXTENSIONS = {
    ".pdf": "application/pdf",
    ".txt": "text/plain",
    ".md": "text/markdown",
    ".html": "text/html",
    ".htm": "text/html",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".json": "application/stix+json",
}
STIX_TEXT_TYPES = ("report", "malware", "intrusion-set", "campaign", "note", "threat-actor", "tool")
MIN_CHARS_PER_PDF_PAGE = 20     # below this on average the PDF is treated as scanned
MAX_OCR_PAGES = 40              # OCR is slow on CPU (about 1-3 s per page)
OCR_DPI = 200

_ocr_engine = None


def ocr_available() -> bool:
    try:
        import rapidocr_onnxruntime  # noqa: F401
        return True
    except ImportError:
        return False


def ocr_page(page) -> list[str]:
    """OCR one PDF page; returns text lines in reading order (top to bottom, left to right)."""
    global _ocr_engine
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR

    if _ocr_engine is None:
        _ocr_engine = RapidOCR()
    pix = page.get_pixmap(dpi=OCR_DPI, colorspace="rgb", alpha=False)
    image = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.height, pix.width, 3)
    boxes, _ = _ocr_engine(image)
    if not boxes:
        return []
    boxes.sort(key=lambda b: (round(b[0][0][1] / 12), b[0][0][0]))       # line bands, then x
    return [text for _, text, score in boxes if float(score) >= 0.5]


def ocr_paragraphs(lines: list[str]) -> list[str]:
    """Join OCR lines into paragraphs: a line ending in . ! ? : closes the paragraph."""
    paragraphs, current = [], []
    for line in lines:
        current.append(line)
        if line.rstrip().endswith((".", "!", "?", ":")):
            paragraphs.append(" ".join(current))
            current = []
    if current:
        paragraphs.append(" ".join(current))
    return paragraphs


class DocumentError(Exception):
    """A file that cannot be turned into text. `code` is the API error code."""

    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code


@dataclass
class Document:
    text: str
    media_type: str
    pages: int | None = None
    page_starts: list[int] = field(default_factory=list)
    warnings: list[str] = field(default_factory=list)
    declared_ttps: list[str] = field(default_factory=list)

    def page_of(self, offset: int) -> int | None:
        """1-based page number of a character offset (None for files without pages)."""
        if not self.page_starts:
            return None
        page = 0
        for i, start in enumerate(self.page_starts):
            if start <= offset:
                page = i
        return page + 1


# ------------------------------------------------------------------ cleaning

_QUOTES = str.maketrans({"‘": "'", "’": "'", "‚": "'", "‛": "'",
                         "“": '"', "”": '"', "„": '"', "″": '"',
                         "–": "-", "—": "-", "−": "-", " ": " ",
                         "​": "", "﻿": "", "­": ""})
_HYPHEN_BREAK = re.compile(r"(\w)-\n(\w)")
_SPACES = re.compile(r"[ \t\f\v]+")
_BLANK_LINES = re.compile(r"\n\s*\n+")


def clean_text(text: str) -> str:
    """Normalise quotes/dashes, join hyphenated line breaks, tidy whitespace.

    Paragraph breaks (blank lines) are kept. Defanged IOCs such as hxxp:// and
    example[.]com are left exactly as they are.
    """
    text = unicodedata.normalize("NFC", text).translate(_QUOTES)
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = _HYPHEN_BREAK.sub(r"\1\2", text)
    lines = [_SPACES.sub(" ", line).strip() for line in text.split("\n")]
    text = "\n".join(lines)
    return _BLANK_LINES.sub("\n\n", text).strip()


def join_paragraphs(paragraphs) -> str:
    return "\n\n".join(p for p in (clean_text(p) for p in paragraphs) if p)


# ------------------------------------------------------------------ readers

def decode_text(data: bytes) -> tuple[str, list[str]]:
    try:
        return data.decode("utf-8-sig"), []
    except UnicodeDecodeError:
        from charset_normalizer import from_bytes

        best = from_bytes(data).best()
        if best is None:
            raise DocumentError("unsupported_file", "the file is not readable text")
        return str(best), [f"text was not UTF-8; decoded as {best.encoding}"]


def read_pdf(data: bytes, ocr: bool = True) -> Document:
    import pymupdf

    try:
        pdf = pymupdf.open(stream=data, filetype="pdf")
    except Exception as exc:  # corrupt or encrypted PDFs
        raise DocumentError("unsupported_file", f"the PDF could not be opened ({exc.__class__.__name__})")
    if pdf.needs_pass:
        raise DocumentError("unsupported_file", "the PDF is password protected")
    pages = []
    for page in pdf:
        blocks = [b[4] for b in page.get_text("blocks", sort=True) if b[6] == 0]   # 0 = text block
        pages.append([clean_text(b) for b in blocks if b.strip()])
    n_pages = len(pages)
    warnings = []
    if sum(len(b) for p in pages for b in p) < MIN_CHARS_PER_PDF_PAGE * max(n_pages, 1):
        if not (ocr and ocr_available()):
            raise DocumentError("no_text_layer", "the PDF has no text layer (it looks scanned) "
                                                 "and OCR is not available on the server")
        if n_pages > MAX_OCR_PAGES:
            raise DocumentError("no_text_layer", f"the PDF is scanned and has {n_pages} pages; "
                                                 f"OCR is limited to {MAX_OCR_PAGES} pages")
        pages = [[clean_text(p) for p in ocr_paragraphs(ocr_page(page))] for page in pdf]
        if sum(len(b) for p in pages for b in p) < MIN_CHARS_PER_PDF_PAGE:
            raise DocumentError("no_text_layer", "the PDF has no text layer and OCR found no readable text")
        warnings.append(f"no text layer: text was read with OCR from {n_pages} page(s); "
                        "OCR can misread words, check the highlighted evidence")
    # Repeated headers/footers: short blocks that appear on more than half of the pages.
    if n_pages >= 3:
        counts = Counter(b for p in pages for b in set(p) if len(b) < 120)
        repeated = {b for b, c in counts.items() if c > n_pages / 2}
        if repeated:
            pages = [[b for b in p if b not in repeated] for p in pages]
            warnings.append(f"{len(repeated)} repeated page header/footer line(s) removed")
    text, starts = "", []
    for blocks in pages:
        page_text = join_paragraphs(blocks)
        if text and page_text:
            text += "\n\n"
        starts.append(len(text))
        text += page_text
    return Document(text, EXTENSIONS[".pdf"], pages=n_pages, page_starts=starts, warnings=warnings)


def read_plain(data: bytes, markdown: bool = False) -> Document:
    text, warnings = decode_text(data)
    if "\x00" in text:
        raise DocumentError("unsupported_file", "the file looks binary, not text")
    if markdown:
        text = re.sub(r"```.*?```", "\n\n", text, flags=re.S)           # code blocks are not prose
        text = re.sub(r"^\s{0,3}#{1,6}\s*", "", text, flags=re.M)        # heading markers
        text = re.sub(r"^\s*[-*+]\s+", "", text, flags=re.M)             # bullets
        text = re.sub(r"!?\[([^\]]*)\]\([^)]*\)", r"\1", text)          # links / images -> label
        text = re.sub(r"(\*\*|__|`)", "", text)
    media = EXTENSIONS[".md" if markdown else ".txt"]
    return Document(clean_text(text), media, warnings=warnings)


def read_html(data: bytes) -> Document:
    from bs4 import BeautifulSoup

    raw, warnings = decode_text(data)
    soup = BeautifulSoup(raw, "html.parser")
    for tag in soup(["script", "style", "noscript", "nav", "header", "footer", "aside",
                     "form", "svg", "iframe", "template"]):
        tag.decompose()
    root = soup.find("article") or soup.find("main") or soup.body or soup
    lines = [line for line in root.get_text("\n").split("\n") if line.strip()]
    return Document(join_paragraphs(lines), EXTENSIONS[".html"], warnings=warnings)


def read_docx(data: bytes) -> Document:
    import docx

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentError("unsupported_file", f"the DOCX could not be opened ({exc.__class__.__name__})")
    paragraphs = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            paragraphs.append(" | ".join(cell.text for cell in row.cells))
    return Document(join_paragraphs(paragraphs), EXTENSIONS[".docx"])


def read_stix(data: bytes) -> Document:
    try:
        bundle = json.loads(decode_text(data)[0])
    except (json.JSONDecodeError, DocumentError):
        raise DocumentError("unsupported_file", "the .json file is not valid JSON")
    if not isinstance(bundle, dict) or bundle.get("type") != "bundle":
        raise DocumentError("unsupported_file", "the .json file is not a STIX 2.1 bundle")
    paragraphs, declared = [], set()
    for obj in bundle.get("objects", []):
        kind = obj.get("type")
        if kind in STIX_TEXT_TYPES:
            for key in ("name", "description", "abstract", "content"):
                value = obj.get(key)
                if isinstance(value, str) and value.strip():
                    paragraphs.append(value)
        elif kind == "attack-pattern":
            for ref in obj.get("external_references", []):
                if ref.get("source_name") == "mitre-attack" and ref.get("external_id"):
                    declared.add(ref["external_id"])
    return Document(join_paragraphs(paragraphs), EXTENSIONS[".json"],
                    declared_ttps=sorted(declared))


def extract_text(data: bytes, filename: str) -> Document:
    """Dispatch on the (already validated) extension. Raises DocumentError."""
    ext = Path(filename).suffix.lower()
    if ext == ".pdf":
        doc = read_pdf(data)
    elif ext in (".txt", ".md"):
        doc = read_plain(data, markdown=ext == ".md")
    elif ext in (".html", ".htm"):
        doc = read_html(data)
    elif ext == ".docx":
        doc = read_docx(data)
    elif ext == ".json":
        doc = read_stix(data)
    else:
        raise DocumentError("unsupported_file", f"unsupported file type {ext!r}")
    if not doc.text.strip():
        raise DocumentError("no_text", "no readable text was found in the file")
    return doc


def from_text(text: str) -> Document:
    """Pasted text."""
    return Document(clean_text(text), "text/plain")
